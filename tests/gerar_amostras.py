# -*- coding: utf-8 -*-
"""Gera as amostras SINTÉTICAS usadas pelos testes (dados fictícios).

Uso em desenvolvimento apenas; as amostras geradas já estão versionadas em
tests/amostras/. Requer: python-docx, python-pptx, openpyxl, reportlab,
LibreOffice (soffice) e qpdf no PATH.

    python tests/gerar_amostras.py
"""
import csv
import io
import json
import os
import random
import shutil
import subprocess
import sys
import tempfile
import zipfile
from email.message import EmailMessage

AQUI = os.path.dirname(os.path.abspath(__file__))
DESTINO = os.path.join(AQUI, 'amostras')
sys.path.insert(0, os.path.join(AQUI, '..', '_analisador'))
from extdados.detectores import cpf_valido, dv_luhn  # noqa: E402

random.seed(20261020)


def cpf(base9):
    for dv in range(100):
        d = base9 + '%02d' % dv
        if cpf_valido(d):
            return d
    raise ValueError(base9)


def fmt_cpf(d):
    return '%s.%s.%s-%s' % (d[:3], d[3:6], d[6:9], d[9:])


def imei(tac8, serie6):
    corpo = tac8 + serie6
    return corpo + dv_luhn(corpo)


CPF_DECL = cpf('529982247')          # 529.982.247-25
CPF_INV = cpf('111444777')           # 111.444.777-35
CPF_TAB1 = cpf('935411347')
CPF_ZERO = cpf('070680938')          # começa com 0 (testa perda do zero no Excel)
IMEI_1 = imei('35209900', '176148')  # 352099001761481
IMEI_2 = imei('86453904', '123456')
IMEI_3 = imei('35693803', '564380')

GABARITO = {
    'email': sorted(['maria.santos@exemplo.com.br', 'fulano.teste@mailinator.com',
                     'pedro.alves@exemplo.org', 'luana.lins@teste.com.br']),
    'telefone': sorted(['+5581987654321', '+5581998765432', '+5581991234567', '+558133334444',
                        '+5581987650001', '+5511912345678']),
    'cpf': sorted([CPF_DECL, CPF_INV, CPF_TAB1, CPF_ZERO, 'M:***456789**']),
    'imei': sorted([IMEI_1[:14], IMEI_2[:14], IMEI_3[:14]]),
    'nome': sorted(['MARIA APARECIDA DOS SANTOS', 'JOSE CARLOS KOWALSKI', 'ANA PAULA CAVALCANTI DE LIMA',
                    'PEDRO HENRIQUE ALVES', 'LUANA FERREIRA LINS']),
}

PARAGRAFOS = [
    'POLÍCIA CIVIL — RELATÓRIO DE TESTE (DADOS FICTÍCIOS)',
    'Declarante: MARIA APARECIDA DOS SANTOS, CPF %s, telefone (81) 98765-4321, '
    'e-mail maria.santos@exemplo.com.br.' % fmt_cpf(CPF_DECL),
    'Investigado: JOSÉ CARLOS KOWALSKI, CPF %s, celular +55 81 99876-5432, WhatsApp 5581991234567.' % CPF_INV,
    'Aparelho apreendido IMEI %s-%s-%s-%s e segundo aparelho IMEI: %s.' % (
        IMEI_1[:2], IMEI_1[2:8], IMEI_1[8:14], IMEI_1[14], IMEI_2),
    'Contato alternativo: fulano.teste@mailinator.com; recado no fixo (81) 3333-4444.',
    'PIX recebido de ***.456.789-** no dia 15/10/2026, valor R$ 1.250,00.',
    'Testemunha: Ana Paula Cavalcanti de Lima, residente na Rua Maria Quitéria, 120, Recife/PE.',
    'Protocolo 2026101912345 e chave 26201012345678000190550010000123451000123456 (não extrair).',
]
CABECALHO = ['Nome', 'CPF', 'Telefone', 'IMEI', 'E-mail']
LINHAS = [
    ['Pedro Henrique Alves', fmt_cpf(CPF_TAB1), '81987650001', IMEI_3, 'pedro.alves@exemplo.org'],
    ['Luana Ferreira Lins', CPF_ZERO, '(11) 91234-5678', IMEI_1[:14] + '0', 'luana.lins@teste.com.br'],
]


def soffice(origem, formato, pasta):
    subprocess.run(['soffice', '--headless', '--convert-to', formato, '--outdir', pasta, origem],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)


def main():
    os.makedirs(DESTINO, exist_ok=True)
    tmp = tempfile.mkdtemp()
    try:
        _gerar(tmp)
        _gerar_ocr(tmp)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    with open(os.path.join(DESTINO, 'gabarito.json'), 'w', encoding='utf-8') as f:
        json.dump(GABARITO, f, ensure_ascii=False, indent=1)
    print('amostras em', DESTINO)


def _gerar(tmp):
    import docx
    import openpyxl
    from pptx import Presentation
    from pptx.util import Inches
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    # ---------------- DOCX
    d = docx.Document()
    d.core_properties.author = 'Equipe de Testes'
    for p in PARAGRAFOS:
        d.add_paragraph(p)
    t = d.add_table(rows=1 + len(LINHAS), cols=len(CABECALHO))
    for i, row in enumerate([CABECALHO] + LINHAS):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    d.sections[0].header.paragraphs[0].text = 'Documento sintético para testes'
    caminho_docx = os.path.join(DESTINO, 'amostra.docx')
    d.save(caminho_docx)
    for fmt in ('pdf', 'doc', 'odt', 'rtf'):
        soffice(caminho_docx, fmt, DESTINO)

    # ---------------- XLSX (números como número: CPF perde o zero à esquerda)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Dados'
    ws.append(['Relatório sintético de testes'])
    ws.append(CABECALHO)
    for linha in LINHAS:
        nova = list(linha)
        for k in (1, 2, 3):
            dig = ''.join(ch for ch in nova[k] if ch.isdigit())
            if nova[k] == dig:
                nova[k] = int(dig)
        ws.append(nova)
    ws2 = wb.create_sheet('Narrativa')
    for p in PARAGRAFOS:
        ws2.append([p])
    caminho_xlsx = os.path.join(DESTINO, 'amostra.xlsx')
    wb.save(caminho_xlsx)
    for fmt in ('xls', 'ods'):
        soffice(caminho_xlsx, fmt, DESTINO)

    # ---------------- CSV (; e Windows-1252), TXT, JSON, HTML
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=';')
    w.writerow(CABECALHO)
    w.writerows(LINHAS)
    with open(os.path.join(DESTINO, 'amostra.csv'), 'w', encoding='cp1252', newline='') as f:
        f.write(buf.getvalue())
    with open(os.path.join(DESTINO, 'amostra.txt'), 'w', encoding='utf-8') as f:
        f.write('\n'.join(PARAGRAFOS) + '\n\n' + '\n'.join('\t'.join(r) for r in [CABECALHO] + LINHAS) + '\n')
    with open(os.path.join(DESTINO, 'amostra.json'), 'w', encoding='utf-8') as f:
        json.dump({'relato': PARAGRAFOS,
                   'pessoas': [dict(zip(['nome', 'cpf', 'telefone', 'imei', 'email'], r)) for r in LINHAS]},
                  f, ensure_ascii=False, indent=1)
    html = ['<!DOCTYPE html><html><head><meta charset="utf-8"><title>Amostra</title></head><body>']
    html += ['<p>%s</p>' % p for p in PARAGRAFOS]
    html.append('<table border="1"><tr>%s</tr>' % ''.join('<th>%s</th>' % c for c in CABECALHO))
    for r in LINHAS:
        html.append('<tr>%s</tr>' % ''.join('<td>%s</td>' % c for c in r))
    html.append('</table><p><a href="mailto:maria.santos@exemplo.com.br">fale conosco</a></p></body></html>')
    with open(os.path.join(DESTINO, 'amostra.html'), 'w', encoding='utf-8') as f:
        f.write('\n'.join(html))
    # ".xls" que na verdade é HTML (exportação típica de sistemas bancários)
    with open(os.path.join(DESTINO, 'amostra_html_disfarcado.xls'), 'w', encoding='cp1252') as f:
        f.write('<html><body><table><tr>%s</tr>%s</table></body></html>' % (
            ''.join('<td>%s</td>' % c for c in CABECALHO),
            ''.join('<tr>%s</tr>' % ''.join('<td>%s</td>' % c for c in r) for r in LINHAS)))

    # ---------------- PPTX
    pr = Presentation()
    s = pr.slides.add_slide(pr.slide_layouts[5])
    s.shapes.title.text = 'Amostra sintética'
    tb = s.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5)).text_frame
    tb.text = PARAGRAFOS[1]
    for p in PARAGRAFOS[2:]:
        tb.add_paragraph().text = p
    pr.save(os.path.join(DESTINO, 'amostra.pptx'))

    # ---------------- PDF com fonte padrão (sem /Widths) e texto posicionado palavra a palavra
    c = canvas.Canvas(os.path.join(DESTINO, 'amostra_reportlab.pdf'), pagesize=A4)
    y = 800
    for p in PARAGRAFOS:
        x = 40
        for palavra in p.split(' '):
            c.setFont('Helvetica', 9)
            c.drawString(x, y, palavra)
            x += c.stringWidth(palavra + ' ', 'Helvetica', 9)
            if x > 520:
                x, y = 40, y - 12
        y -= 16
    for r in [CABECALHO] + LINHAS:
        x = 40
        for v in r:
            c.drawString(x, y, v)
            x += 110
        y -= 14
    c.save()

    # ---------------- PDFs criptografados (senha de abertura vazia) e com senha
    base = os.path.join(DESTINO, 'amostra.pdf')
    for nome, args in (('amostra_rc4.pdf', ['128', '--use-aes=n']),
                       ('amostra_aes128.pdf', ['128', '--use-aes=y']),
                       ('amostra_aes256.pdf', ['256'])):
        subprocess.run(['qpdf', '--allow-weak-crypto', '--encrypt', '', 'dono123'] + args + ['--', base, os.path.join(DESTINO, nome)],
                       check=True)
    subprocess.run(['qpdf', '--encrypt', 'abrir123', 'dono123', '256', '--', base,
                    os.path.join(DESTINO, 'amostra_com_senha.pdf')], check=True)
    subprocess.run(['qpdf', '--object-streams=generate', '--compress-streams=y', base,
                    os.path.join(DESTINO, 'amostra_objstm.pdf')], check=True)

    # ---------------- EML com anexo CSV e ZIP com HTML + TXT
    m = EmailMessage()
    m['From'] = 'Maria Aparecida dos Santos <maria.santos@exemplo.com.br>'
    m['To'] = 'fulano.teste@mailinator.com'
    m['Subject'] = 'Dados do caso (fictício)'
    m.set_content('\n'.join(PARAGRAFOS))
    with open(os.path.join(DESTINO, 'amostra.csv'), 'rb') as f:
        m.add_attachment(f.read(), maintype='text', subtype='csv', filename='planilha anexa.csv')
    with open(os.path.join(DESTINO, 'amostra.eml'), 'wb') as f:
        f.write(bytes(m))
    with zipfile.ZipFile(os.path.join(DESTINO, 'amostra.zip'), 'w', zipfile.ZIP_DEFLATED) as z:
        z.write(os.path.join(DESTINO, 'amostra.html'), 'pasta/registros.html')
        z.write(os.path.join(DESTINO, 'amostra.txt'), 'relato.txt')


# ----------------------------------------------------------------------------- OCR
OCR_LINHAS = ['TERMO DE DECLARACOES (DOCUMENTO DIGITALIZADO FICTICIO)',
              'Declarante: JOAO BATISTA DE OLIVEIRA, CPF %s,' % fmt_cpf(CPF_DECL),
              'telefone (81) 98877-6655, e-mail joao.batista@exemplo.com.br,',
              'aparelho IMEI %s-%s-%s-%s.' % (IMEI_1[:2], IMEI_1[2:8], IMEI_1[8:14], IMEI_1[14])]
GABARITO_OCR = {'cpf': CPF_DECL, 'telefone': '+5581988776655', 'imei': IMEI_1[:14],
                'email': 'joao.batista@exemplo.com.br'}


def _pagina_escaneada(largura=2480, altura=1100):
    import glob
    from PIL import Image, ImageDraw, ImageFont
    fonte = ImageFont.truetype(glob.glob('/usr/share/fonts/**/DejaVuSans.ttf', recursive=True)[0], 40)
    img = Image.new('L', (largura, altura), 255)
    d = ImageDraw.Draw(img)
    for i, linha in enumerate(OCR_LINHAS):
        d.text((140, 140 + i * 100), linha, fill=0, font=fonte)
    return img


def _gerar_ocr(tmp):
    """Amostras para o OCR: imagem solta, PDFs digitalizados em vários formatos de imagem,
    PDF já pesquisável (com camada de texto invisível) e print colado num DOCX."""
    import img2pdf
    import docx
    from docx.shared import Inches
    img = _pagina_escaneada()
    png = os.path.join(DESTINO, 'ocr_print.png')
    img.save(png, dpi=(300, 300))
    jpg = os.path.join(tmp, 'pagina.jpg')
    img.save(jpg, quality=80, dpi=(300, 300))
    cinza = os.path.join(tmp, 'pagina_cinza.png')
    img.save(cinza, dpi=(300, 300))
    paleta = os.path.join(tmp, 'pagina_paleta.png')
    img.convert('RGB').quantize(16).save(paleta, dpi=(300, 300))
    pb = os.path.join(tmp, 'pagina_pb.png')
    img.convert('1').save(pb, dpi=(300, 300))
    girada = os.path.join(tmp, 'pagina_girada.png')
    img.rotate(90, expand=True).save(girada, dpi=(300, 300))
    for origem, destino in ((jpg, 'ocr_scan_jpeg.pdf'), (cinza, 'ocr_scan_cinza.pdf'),
                            (paleta, 'ocr_scan_paleta.pdf'), (girada, 'ocr_scan_girado.pdf')):
        with open(os.path.join(DESTINO, destino), 'wb') as f:
            f.write(img2pdf.convert(origem))
    subprocess.run(['convert', pb, '-compress', 'Group4', os.path.join(DESTINO, 'ocr_scan_ccitt.pdf')],
                   check=True)
    # PDF pesquisável (scanner com OCR embutido): texto invisível sobre a imagem
    subprocess.run(['tesseract', cinza, os.path.join(DESTINO, 'ocr_scan_pesquisavel'), '-l', 'eng', 'pdf'],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    # página com texto digitado + print colado (relatório exportado para PDF)
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(os.path.join(DESTINO, 'ocr_misto.pdf'), pagesize=A4)
    c.setFont('Helvetica', 10)
    for i, linha in enumerate(PARAGRAFOS[:3]):
        c.drawString(40, 800 - i * 16, linha[:110])
    c.drawImage(png, 40, 420, width=515, height=515 * 1100 / 2480.0)
    c.save()
    d = docx.Document()
    d.add_paragraph('Relatório com print colado (fictício). Segue imagem:')
    d.add_picture(png, width=Inches(6))
    d.save(os.path.join(DESTINO, 'ocr_print_no_docx.docx'))
    with open(os.path.join(DESTINO, 'gabarito_ocr.json'), 'w', encoding='utf-8') as f:
        json.dump(GABARITO_OCR, f, ensure_ascii=False, indent=1)


if __name__ == '__main__':
    if sys.argv[1:] == ['--so-ocr']:
        _t = tempfile.mkdtemp()
        try:
            _gerar_ocr(_t)
        finally:
            shutil.rmtree(_t, ignore_errors=True)
    else:
        main()
